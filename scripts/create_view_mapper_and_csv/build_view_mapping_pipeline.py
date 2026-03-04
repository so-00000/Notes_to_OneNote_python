from __future__ import annotations

import argparse
import csv
import json
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Sequence

BASE_DIR = Path(__file__).resolve().parent
SOURCE_ROOT = BASE_DIR / "source_view"
MAPPING_ROOT = (BASE_DIR / "../../main/resources/mapping").resolve()
COMMON_MAPPING_PATH = MAPPING_ROOT / "cmn_mapping.json"

COMMON_COLUMNS: List[Dict[str, str]] = [
    {"key": "source_id", "meaning": "一意キー（推奨: replica_id + '_' + unid）", "source": "computed"},
    {"key": "replica_id", "meaning": "Notes DB識別子", "source": "Notes/DXL/設定"},
    {"key": "unid", "meaning": "Notes文書UNID", "source": "Notes/DXL"},
    {"key": "form", "meaning": "フォーム名", "source": "Notes item Form"},
    {"key": "title", "meaning": "ページタイトル", "source": "Notes item or ルール"},
    {"key": "doc_date", "meaning": "並び替え基準日", "source": "Notes item or ルール"},
    {"key": "onenote_page_id", "meaning": "OneNoteページID", "source": "Graphレスポンス"},
    {"key": "onenote_web_url", "meaning": "Web用URL", "source": "Graphレスポンス links.oneNoteWebUrl.href"},
    {"key": "onenote_client_url", "meaning": "クライアント用URL", "source": "Graphレスポンス links.oneNoteClientUrl.href"},
    {"key": "status", "meaning": "処理状態（SUCCESS/FAILED/SKIPPED等）", "source": "処理ロジック"},
    {"key": "error_message", "meaning": "失敗時のエラーメッセージ", "source": "例外/レスポンス"},
    {"key": "processed_at", "meaning": "処理日時", "source": "処理ロジック"},
]


@dataclass
class ViewColumn:
    column_index: int
    column_title: str
    item_name: str
    formula: str
    is_category: bool
    is_hidden: bool
    output_key: str


@dataclass
class ViewPaths:
    view_folder: str
    source_dir: Path
    mapping_path: Path
    master_csv_path: Path
    excel_path: Path


def _read_view_csv(csv_path: Path) -> List[Dict[str, str]]:
    for enc in ("utf-8-sig", "cp932", "utf-8"):
        try:
            with csv_path.open("r", encoding=enc, newline="") as f:
                return list(csv.DictReader(f))
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("unknown", b"", 0, 1, f"could not decode: {csv_path}")


def _slugify(value: str) -> str:
    value = value.strip().lower()
    value = re.sub(r"\s+", "_", value)
    value = re.sub(r"[^a-z0-9_]", "_", value)
    value = re.sub(r"_+", "_", value).strip("_")
    return value or "col"


def _bool_text(v: str) -> bool:
    return str(v).strip() in {"1", "true", "True", "TRUE"}


def _build_view_columns(rows: Sequence[Dict[str, str]]) -> List[ViewColumn]:
    cols: List[ViewColumn] = []
    seen_keys: Dict[str, int] = {}
    for row in sorted(rows, key=lambda r: int(r.get("ColumnIndex") or 0)):
        title = (row.get("ColumnTitle") or "").strip()
        item_name = (row.get("ItemName") or "").strip()
        base_key = _slugify(item_name or title or f"col_{row.get('ColumnIndex', '')}")
        seen_keys[base_key] = seen_keys.get(base_key, 0) + 1
        output_key = base_key if seen_keys[base_key] == 1 else f"{base_key}_{seen_keys[base_key]}"
        cols.append(
            ViewColumn(
                column_index=int(row.get("ColumnIndex") or 0),
                column_title=title,
                item_name=item_name,
                formula=(row.get("Formula") or "").strip(),
                is_category=_bool_text(row.get("IsCategory", "")),
                is_hidden=_bool_text(row.get("IsHidden", "")),
                output_key=output_key,
            )
        )
    return cols


def _list_view_folders() -> List[str]:
    if not SOURCE_ROOT.exists():
        return []
    return sorted([p.name for p in SOURCE_ROOT.iterdir() if p.is_dir()])


def _select_view_folder(cli_view: str | None) -> str:
    folders = _list_view_folders()
    if not folders:
        raise FileNotFoundError(f"Viewフォルダが見つかりません: {SOURCE_ROOT}")

    if cli_view:
        if cli_view not in folders:
            raise ValueError(f"指定Viewフォルダが存在しません: {cli_view}")
        return cli_view

    if len(folders) == 1:
        print(f"対象View: {folders[0]} (唯一のフォルダ)")
        return folders[0]

    print("対象Viewフォルダを選択してください:")
    for i, name in enumerate(folders, start=1):
        print(f"  {i}. {name}")

    while True:
        answer = input("番号を入力: ").strip()
        if answer.isdigit():
            idx = int(answer)
            if 1 <= idx <= len(folders):
                return folders[idx - 1]
        print("無効な入力です。表示された番号を入力してください。")


def _resolve_paths(view_folder: str) -> ViewPaths:
    source_dir = SOURCE_ROOT / view_folder

    return ViewPaths(
        view_folder=view_folder,
        source_dir=source_dir,
        mapping_path=MAPPING_ROOT / view_folder / "mapping.json",
        master_csv_path=MAPPING_ROOT / view_folder / "migration_master.csv",
        excel_path=MAPPING_ROOT / view_folder / "migration_master.xlsx",
    )


def generate_mapping(paths: ViewPaths) -> Dict[str, object]:
    csv_files = sorted(paths.source_dir.glob("*.csv"))
    if not csv_files:
        raise FileNotFoundError(f"CSVが見つかりません: {paths.source_dir}")

    views: Dict[str, Dict[str, object]] = {}
    for csv_file in csv_files:
        rows = _read_view_csv(csv_file)
        if not rows:
            continue
        view_name = (rows[0].get("ViewName") or csv_file.stem).strip()
        cols = _build_view_columns(rows)
        views[view_name] = {
            "source_csv": csv_file.name,
            "view_columns": [
                {
                    "column_index": c.column_index,
                    "column_title": c.column_title,
                    "item_name": c.item_name,
                    "formula": c.formula,
                    "is_category": c.is_category,
                    "is_hidden": c.is_hidden,
                    "output_key": c.output_key,
                }
                for c in cols
            ],
        }

    mapping = {
        "version": 1,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "view_folder": paths.view_folder,
        "common_columns": COMMON_COLUMNS,
        "views": views,
    }
    paths.mapping_path.parent.mkdir(parents=True, exist_ok=True)
    with paths.mapping_path.open("w", encoding="utf-8") as f:
        json.dump(mapping, f, ensure_ascii=False, indent=2)
    return mapping


def write_common_mapping() -> None:
    common = {
        "version": 1,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "common_columns": COMMON_COLUMNS,
    }
    COMMON_MAPPING_PATH.parent.mkdir(parents=True, exist_ok=True)
    with COMMON_MAPPING_PATH.open("w", encoding="utf-8") as f:
        json.dump(common, f, ensure_ascii=False, indent=2)


def _collect_master_headers(mapping: Dict[str, object]) -> List[str]:
    common = [c["key"] for c in mapping["common_columns"]]
    headers: List[str] = ["view_name", *common]
    seen = set(headers)
    for view in mapping["views"].values():
        for col in view["view_columns"]:
            key = col["output_key"]
            if key not in seen:
                headers.append(key)
                seen.add(key)
    return headers


def sync_master_csv(paths: ViewPaths) -> List[str]:
    with paths.mapping_path.open("r", encoding="utf-8") as f:
        mapping = json.load(f)

    expected_headers = _collect_master_headers(mapping)
    existing_rows: List[Dict[str, str]] = []
    if paths.master_csv_path.exists():
        for enc in ("utf-8-sig", "cp932", "utf-8"):
            try:
                with paths.master_csv_path.open("r", encoding=enc, newline="") as f:
                    existing_rows = list(csv.DictReader(f))
                break
            except UnicodeDecodeError:
                continue

    normalized_rows: List[Dict[str, str]] = []
    for row in existing_rows:
        normalized_rows.append({h: row.get(h, "") for h in expected_headers})

    paths.master_csv_path.parent.mkdir(parents=True, exist_ok=True)
    with paths.master_csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=expected_headers)
        writer.writeheader()
        writer.writerows(normalized_rows)
    return expected_headers


def export_excel(paths: ViewPaths) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:
        raise RuntimeError("Excel出力には openpyxl が必要です。`pip install openpyxl` を実行してください。") from exc

    rows: List[Dict[str, str]] = []
    headers: List[str] = []
    for enc in ("utf-8-sig", "cp932", "utf-8"):
        try:
            with paths.master_csv_path.open("r", encoding=enc, newline="") as f:
                reader = csv.DictReader(f)
                headers = reader.fieldnames or []
                rows = list(reader)
            break
        except UnicodeDecodeError:
            continue

    if not headers:
        raise ValueError(f"成果CSVのヘッダが取得できません: {paths.master_csv_path}")

    with paths.mapping_path.open("r", encoding="utf-8") as f:
        mapping = json.load(f)

    common_meaning_by_key = {c["key"]: c["meaning"] for c in mapping.get("common_columns", [])}
    view_column_title_by_key: Dict[str, str] = {}
    for view_def in mapping.get("views", {}).values():
        for col in view_def.get("view_columns", []):
            key = col.get("output_key", "")
            if key and key not in view_column_title_by_key:
                title = (col.get("column_title") or "").strip()
                item_name = (col.get("item_name") or "").strip()
                view_column_title_by_key[key] = title or item_name or key

    logical_headers: List[str] = []
    for h in headers:
        if h == "view_name":
            logical_headers.append("View名")
            continue
        if h in common_meaning_by_key:
            logical_headers.append(common_meaning_by_key[h])
            continue
        logical_headers.append(view_column_title_by_key.get(h, h))

    wb = Workbook()
    ws = wb.active
    ws.title = "migration_result"
    ws.append(logical_headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])

    font = Font(name="Meiryo UI")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.font = font

    paths.excel_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(paths.excel_path)


def build_all(paths: ViewPaths) -> None:
    write_common_mapping()
    mapping = generate_mapping(paths)
    headers = sync_master_csv(paths)
    export_excel(paths)
    print(f"view_folder: {paths.view_folder}")
    print(f"mapping.json を生成: {paths.mapping_path}")
    print(f"共通 mapping を生成: {COMMON_MAPPING_PATH}")
    print(f"成果CSVを更新: {paths.master_csv_path} (columns={len(headers)}, views={len(mapping['views'])})")
    print(f"Excelを出力: {paths.excel_path}")


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Viewフォルダ単位で mapping.json / 成果CSV / Excel を生成するコマンド"
    )

    subparsers = parser.add_subparsers(dest="subcommand")
    parser.add_argument("--view", type=str, default=None, help="source_view 配下の対象Viewフォルダ名")
    build_parser = subparsers.add_parser("build", help="mapping.json -> 成果CSV更新 -> Excel出力を一括実行")
    gen_parser = subparsers.add_parser("generate-mapping", help="mapping.jsonのみ生成")
    sync_parser = subparsers.add_parser("sync-master-csv", help="mapping.jsonを使って成果CSVの列を同期")
    excel_parser = subparsers.add_parser("export-excel", help="成果CSVからExcel出力")
    subparsers.add_parser("list-views", help="選択可能なViewフォルダ一覧を表示")

    for sub in (build_parser, gen_parser, sync_parser, excel_parser):
        sub.add_argument("--view", type=str, default=None, help="source_view 配下の対象Viewフォルダ名")
    return parser


def main() -> None:
    parser = _build_parser()
    args = parser.parse_args()

    cmd = args.subcommand or "build"
    if cmd == "list-views":
        for name in _list_view_folders():
            print(name)
        return

    selected_view = getattr(args, "view", None)
    view_folder = _select_view_folder(selected_view)
    paths = _resolve_paths(view_folder)

    if cmd == "build":
        build_all(paths)
        return
    if cmd == "generate-mapping":
        write_common_mapping()
        mapping = generate_mapping(paths)
        print(f"view_folder: {paths.view_folder}")
        print(f"mapping.json を生成: {paths.mapping_path} (views={len(mapping['views'])})")
        print(f"共通 mapping を生成: {COMMON_MAPPING_PATH}")
        return
    if cmd == "sync-master-csv":
        headers = sync_master_csv(paths)
        print(f"view_folder: {paths.view_folder}")
        print(f"成果CSVを更新: {paths.master_csv_path} (columns={len(headers)})")
        return
    if cmd == "export-excel":
        export_excel(paths)
        print(f"view_folder: {paths.view_folder}")
        print(f"Excelを出力: {paths.excel_path}")
        return
    parser.error(f"未知のサブコマンド: {cmd}")


if __name__ == "__main__":
    main()
